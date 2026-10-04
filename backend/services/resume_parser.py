from io import BytesIO

import pymupdf
from docx import Document


MAX_EXTRACTED_TEXT_CHARS = 100_000


class ResumeExtractionError(ValueError):
    pass


def extract_resume_text(file_bytes, content_type):
    if content_type == "application/pdf":
        text = _extract_pdf_text(file_bytes)
    elif (
        content_type
        == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    ):
        text = _extract_docx_text(file_bytes)
    else:
        raise ResumeExtractionError("Use a PDF or DOCX file for your CV.")

    text = text.strip()
    if not text:
        raise ResumeExtractionError(
            "No selectable text was found. Scanned CVs aren't supported yet."
        )
    if len(text) > MAX_EXTRACTED_TEXT_CHARS:
        raise ResumeExtractionError("The extracted CV text is too long to process.")
    return text


def _extract_pdf_text(file_bytes):
    try:
        with pymupdf.open(stream=file_bytes, filetype="pdf") as document:
            if document.needs_pass:
                raise ResumeExtractionError("Password-protected PDFs aren't supported.")
            return "\n".join(page.get_text("text") for page in document)
    except (pymupdf.FileDataError, RuntimeError) as exc:
        raise ResumeExtractionError("The PDF file could not be read.") from exc


def _extract_docx_text(file_bytes):
    try:
        document = Document(BytesIO(file_bytes))
    except Exception as exc:
        raise ResumeExtractionError("The DOCX file could not be read.") from exc

    blocks = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            blocks.append("\t".join(cell.text for cell in row.cells))
    return "\n".join(blocks)